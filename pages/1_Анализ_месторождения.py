"""
pages/1_Анализ_месторождения.py — анализ "сырых" лабораторных PVT-данных по
месторождению: загрузка таблицы отборов проб (много скважин, горизонты,
блоки), гибкое сопоставление колонок, фильтры и построение зависимостей
между свойствами (Rs vs Pb, Bo vs Pb, вязкость vs Pb, свойство vs глубина
и т.д.) — как по всему массиву, так и с группировкой по горизонту/блоку.

Формат исходной таблицы заранее не фиксирован: шапка может занимать 1-3
строки, колонки "Горизонт"/"блок" могут отсутствовать вовсе — тогда просто
не показываются соответствующие фильтры и группировки.
"""

from __future__ import annotations

import io

import numpy as np
import openpyxl
import pandas as pd
import plotly.graph_objects as go
import streamlit as st

import field_data as fd
import pvt_tuning as pt
import units as u

st.set_page_config(page_title="Анализ месторождения — PVT", layout="wide")
st.title("Анализ лабораторных PVT-данных по месторождению")
st.caption(
    "Загрузите таблицу отборов проб (xlsx/csv) — программа попробует сама "
    "распознать шапку и колонки, но всё можно поправить руками ниже."
)

# ---------------------------------------------------------------------------
# 1. Загрузка файла
# ---------------------------------------------------------------------------
uploaded = st.file_uploader("Файл с лабораторными PVT-данными (xlsx или csv)",
                             type=["xlsx", "xls", "csv"])

if uploaded is None:
    st.info("Загрузите файл, чтобы начать анализ.")
    st.stop()

is_excel = uploaded.name.lower().endswith((".xlsx", ".xls"))

if is_excel:
    wb = openpyxl.load_workbook(uploaded, data_only=True)
    sheet_name = st.selectbox("Лист", options=wb.sheetnames, index=0) \
        if len(wb.sheetnames) > 1 else wb.sheetnames[0]
    ws = wb[sheet_name]
    all_rows = list(ws.iter_rows(values_only=True))
else:
    df_csv_raw = pd.read_csv(uploaded, header=None)
    all_rows = [tuple(r) for r in df_csv_raw.itertuples(index=False, name=None)]

if not all_rows:
    st.error("Файл пуст.")
    st.stop()

# ---------------------------------------------------------------------------
# 2. Определение шапки и начала данных (с автоподсказкой)
# ---------------------------------------------------------------------------
st.subheader("1. Шапка таблицы и начало данных")

guess_header_idx, guess_data_idx = fd.guess_header_and_data_rows(all_rows)

with st.expander("Предпросмотр листа (первые 40 строк)", expanded=False):
    preview_rows = all_rows[:40]
    max_cols = max((len(r) for r in preview_rows), default=0)
    df_preview = pd.DataFrame(
        [list(r) + [None] * (max_cols - len(r)) for r in preview_rows],
        index=[f"Строка {i + 1}" for i in range(len(preview_rows))],
    )
    st.dataframe(df_preview, use_container_width=True)

col_a, col_b = st.columns(2)
with col_a:
    header_start_row = st.number_input(
        "Первая строка шапки (номер строки в файле, с 1)", min_value=1,
        max_value=len(all_rows), value=guess_header_idx + 1,
    )
with col_b:
    data_start_row = st.number_input(
        "Первая строка данных (номер строки в файле, с 1)", min_value=1,
        max_value=len(all_rows), value=guess_data_idx + 1,
    )

if data_start_row <= header_start_row:
    st.error("Строка данных должна быть ниже начала шапки.")
    st.stop()

header_rows = all_rows[header_start_row - 1: data_start_row - 1]
headers = fd.merge_header_rows(header_rows)

data_rows_raw = [r for r in all_rows[data_start_row - 1:] if any(c is not None for c in r)]
n_cols = len(headers)
data_rows = [tuple(r[:n_cols]) + (None,) * max(0, n_cols - len(r)) for r in data_rows_raw]
df_raw = pd.DataFrame(data_rows, columns=headers)

st.success(f"Распознано {len(headers)} колонок и {len(df_raw)} строк данных.")

# ---------------------------------------------------------------------------
# 3. Сопоставление колонок со стандартными полями
# ---------------------------------------------------------------------------
st.subheader("2. Сопоставление колонок")
st.caption(
    "Автоматически предложенное сопоставление — проверьте и при необходимости "
    "исправьте. Поля, для которых нет подходящей колонки, оставьте "
    "«— не использовать —»."
)

suggested = fd.suggest_column_mapping(headers)
options = ["— не использовать —"] + headers

mapping: dict[str, str] = {}
map_cols = st.columns(3)
all_field_keys = [key for key, _label, _aliases in fd.FIELD_ALIASES]
for i, key in enumerate(all_field_keys):
    label = fd.FIELD_LABELS[key]
    default_header = suggested.get(key)
    default_idx = options.index(default_header) if default_header in options else 0
    with map_cols[i % 3]:
        chosen = st.selectbox(label, options=options, index=default_idx, key=f"map_{key}")
    if chosen != "— не использовать —":
        mapping[key] = chosen

if "pb" not in mapping and "p_sample" not in mapping:
    st.warning("Не сопоставлена ни одна колонка давления (Pb или давление отбора) — "
               "большинство типовых зависимостей строить будет не с чем.")

# ---------------------------------------------------------------------------
# 4. Построение чистой таблицы, единицы давления, агрегатные строки
# ---------------------------------------------------------------------------
df_clean = fd.build_clean_dataframe(df_raw, mapping)

st.subheader("3. Единицы давления и очистка данных")
pressure_units = ["МПа", "бар", "psi"]
c1, c2 = st.columns(2)
with c1:
    source_p_unit = st.selectbox("Единица давления В ФАЙЛЕ", options=pressure_units, index=0)
with c2:
    display_p_unit = st.selectbox("Единица давления для отображения/графиков",
                                   options=pressure_units, index=0)


def _to_mpa(val: float, unit: str) -> float:
    if unit == "МПа":
        return val
    if unit == "бар":
        return u.bar_to_mpa(val)
    return u.psi_to_mpa(val)


def _from_mpa(val: float, unit: str) -> float:
    if unit == "МПа":
        return val
    if unit == "бар":
        return u.mpa_to_bar(val)
    return u.mpa_to_psi(val)


for pcol in ("pb", "p_sample"):
    if pcol in df_clean.columns:
        df_clean[pcol] = df_clean[pcol].apply(
            lambda v: _from_mpa(_to_mpa(v, source_p_unit), display_p_unit) if pd.notna(v) else v
        )

qc_col1, qc_col2 = st.columns(2)
with qc_col1:
    exclude_aggregates = st.checkbox(
        f"Исключить агрегатные строки типа «Среднее значение: …» "
        f"(найдено: {int(df_clean['is_aggregate'].sum())})",
        value=True,
    )
with qc_col2:
    n_anomalous = int(df_clean["is_anomalous"].sum())
    exclude_anomalous = st.checkbox(
        f"Исключить пробы, помеченные звёздочкой (*) как непредставительные "
        f"(найдено: {n_anomalous})",
        value=n_anomalous > 0,
        disabled=n_anomalous == 0,
    )

if exclude_aggregates:
    df_clean = df_clean[~df_clean["is_aggregate"]].reset_index(drop=True)
if exclude_anomalous:
    df_clean = df_clean[~df_clean["is_anomalous"]].reset_index(drop=True)

# Кросс-проверка Rs(м3/т) vs Rs(м3/м3) через материальный баланс и флаг
# "P отбора ниже Pb" — не отбрасываем строки, только предупреждаем
cross_check_issues = df_clean[df_clean["rs_cross_check_dev_pct"] > 5.0] \
    if "rs_cross_check_dev_pct" in df_clean.columns else pd.DataFrame()
p_below_pb_issues = df_clean[df_clean.get("p_below_pb", pd.Series(dtype=bool)) == True] \
    if "p_below_pb" in df_clean.columns else pd.DataFrame()

if not cross_check_issues.empty:
    with st.expander(
        f"⚠ {len(cross_check_issues)} проб(ы) с расхождением Rs(м³/т) и Rs(м³/м³) "
        f"более 5% — возможна опечатка в исходных данных", expanded=False,
    ):
        cols_show = [c for c in ("horizon", "well", "rs_m3m3", "rs_m3t",
                                  "rs_cross_check_dev_pct") if c in cross_check_issues.columns]
        st.dataframe(cross_check_issues[cols_show].sort_values(
            "rs_cross_check_dev_pct", ascending=False), use_container_width=True)

if not p_below_pb_issues.empty:
    with st.expander(
        f"ℹ {len(p_below_pb_issues)} проб(ы) с давлением отбора ниже давления "
        f"насыщения — при отборе в пласте уже мог идти двухфазный поток", expanded=False,
    ):
        cols_show = [c for c in ("horizon", "well", "pb", "p_sample")
                     if c in p_below_pb_issues.columns]
        st.dataframe(p_below_pb_issues[cols_show], use_container_width=True)

# ---------------------------------------------------------------------------
# 5. Фильтры (показываются только если соответствующее поле сопоставлено)
# ---------------------------------------------------------------------------
st.subheader("4. Фильтры")
df_filtered = df_clean.copy()

filter_cols = st.columns(4)
group_field_present = {}
for i, key in enumerate(fd.GROUP_FIELDS):
    if key in mapping and df_clean[key].notna().any():
        values = sorted(v for v in df_clean[key].dropna().unique())
        with filter_cols[i % 4]:
            selected = st.multiselect(fd.FIELD_LABELS[key], options=values, default=[])
        if selected:
            df_filtered = df_filtered[df_filtered[key].isin(selected)]
        group_field_present[key] = True
    else:
        group_field_present[key] = False

if not any(group_field_present.values()):
    st.caption("В данных нет колонок горизонта/блока/скважины/вида пробы — "
               "фильтры и группировка по ним недоступны, зависимости строятся по всему массиву.")

st.dataframe(df_filtered, use_container_width=True)
st.caption(f"После фильтров: {len(df_filtered)} строк из {len(df_clean)}.")

# ---------------------------------------------------------------------------
# 6. Построение зависимостей
# ---------------------------------------------------------------------------
st.subheader("5. Зависимости")

available_numeric = [k for k in fd.NUMERIC_FIELDS if k in mapping and df_filtered[k].notna().any()]

if len(available_numeric) < 2:
    st.warning("Недостаточно числовых полей для построения зависимостей "
               "(нужно минимум два сопоставленных численных свойства).")
    st.stop()


def _axis_label(key: str) -> str:
    if key in ("pb", "p_sample"):
        base = "Давление насыщения" if key == "pb" else "Давление отбора"
        return f"{base}, {display_p_unit}"
    unit = fd.FIELD_UNITS.get(key)
    return f"{fd.FIELD_LABELS[key]}, {unit}" if unit else fd.FIELD_LABELS[key]


if "x_field" not in st.session_state:
    st.session_state.x_field = ("pb" if "pb" in available_numeric else available_numeric[0])
if "y_field" not in st.session_state:
    st.session_state.y_field = ("rs_m3m3" if "rs_m3m3" in available_numeric else available_numeric[-1])

st.caption("Быстрые пресеты:")
preset_cols = st.columns(5)
presets = [
    ("Rs vs Pb", "pb", "rs_m3m3"),
    ("Bo vs Pb", "pb", "bo"),
    ("μo vs Pb", "pb", "mu_oil"),
    ("Rs vs Глубина", "depth", "rs_m3m3"),
    ("Bo vs Глубина", "depth", "bo"),
]
for col, (label, xk, yk) in zip(preset_cols, presets):
    with col:
        if st.button(label, disabled=not (xk in available_numeric and yk in available_numeric)):
            st.session_state.x_field = xk
            st.session_state.y_field = yk

c1, c2, c3 = st.columns(3)
with c1:
    x_field = st.selectbox("X (по горизонтали)", options=available_numeric,
                            format_func=_axis_label,
                            index=available_numeric.index(st.session_state.x_field)
                            if st.session_state.x_field in available_numeric else 0)
with c2:
    y_field = st.selectbox("Y (по вертикали)", options=available_numeric,
                            format_func=_axis_label,
                            index=available_numeric.index(st.session_state.y_field)
                            if st.session_state.y_field in available_numeric else 0)
with c3:
    trend_kind_label = st.selectbox("Линия тренда", options=["Линейная", "Степенная", "Без тренда"])
    trend_kind = {"Линейная": "linear", "Степенная": "power", "Без тренда": None}[trend_kind_label]

group_options = ["Без группировки (весь массив/горизонт целиком)"] + \
    [fd.FIELD_LABELS[k] for k in fd.GROUP_FIELDS if group_field_present.get(k)]
group_choice = st.selectbox("Группировка", options=group_options)

group_key = None
if group_choice != group_options[0]:
    for k in fd.GROUP_FIELDS:
        if group_field_present.get(k) and fd.FIELD_LABELS[k] == group_choice:
            group_key = k
            break

plot_df = df_filtered[[x_field, y_field] + ([group_key] if group_key else [])].dropna(
    subset=[x_field, y_field]
)

if plot_df.empty:
    st.warning("Нет точек с одновременно заполненными X и Y для выбранных полей/фильтров.")
    st.stop()

fig = go.Figure()
trend_rows = []

palette = ["#1f77b4", "#ff7f0e", "#2ca02c", "#d62728", "#9467bd", "#8c564b",
           "#e377c2", "#7f7f7f", "#bcbd22", "#17becf"]

if group_key:
    groups = list(plot_df.groupby(group_key))
    for i, (gval, gdf) in enumerate(groups):
        color = palette[i % len(palette)]
        fig.add_trace(go.Scatter(
            x=gdf[x_field], y=gdf[y_field], mode="markers", name=f"{gval} (n={len(gdf)})",
            marker=dict(color=color, size=8),
        ))
        if trend_kind:
            res = fd.fit_trend(gdf[x_field].to_numpy(), gdf[y_field].to_numpy(), kind=trend_kind)
            if res:
                fig.add_trace(go.Scatter(
                    x=res["x_fit"], y=res["y_fit"], mode="lines", name=f"{gval}: тренд",
                    line=dict(color=color, dash="dash"), showlegend=False,
                ))
                trend_rows.append({"Группа": str(gval), "n": len(gdf),
                                    "Уравнение": res["equation"], "R²": round(res["r2"], 4)})
else:
    fig.add_trace(go.Scatter(x=plot_df[x_field], y=plot_df[y_field], mode="markers",
                              name="Данные", marker=dict(color=palette[0], size=8)))
    if trend_kind:
        res = fd.fit_trend(plot_df[x_field].to_numpy(), plot_df[y_field].to_numpy(), kind=trend_kind)
        if res:
            fig.add_trace(go.Scatter(x=res["x_fit"], y=res["y_fit"], mode="lines",
                                      name="Тренд", line=dict(color="red", dash="dash")))
            trend_rows.append({"Группа": "весь массив", "n": len(plot_df),
                                "Уравнение": res["equation"], "R²": round(res["r2"], 4)})

fig.update_layout(
    title=f"{_axis_label(y_field)} vs {_axis_label(x_field)}",
    xaxis_title=_axis_label(x_field), yaxis_title=_axis_label(y_field),
    hovermode="closest", height=600,
)
st.plotly_chart(fig, use_container_width=True)

if trend_rows:
    st.markdown("**Параметры линий тренда**")
    st.dataframe(pd.DataFrame(trend_rows), use_container_width=True)

# ---------------------------------------------------------------------------
# 6. Калибровка (tuning) корреляций Standing под фактические данные
# ---------------------------------------------------------------------------
st.markdown("---")
st.subheader("6. Калибровка корреляций Standing под факт (tuning)")
st.caption(
    "Сравниваем, что предсказала бы корреляция Standing/Beggs-Robinson для "
    "этих же проб (по их Rsb и T), с тем, что реально измерено в лаборатории, "
    "и находим множитель, который сводит модель к факту. Коэффициент можно "
    "сохранить и применить на главной странице (калькуляторе)."
)

required_tuning_fields = ["pb", "rs_m3m3", "bo", "mu_oil", "t_sample_c"]
missing_tuning = [f for f in required_tuning_fields if f not in mapping]

if missing_tuning:
    st.warning("Для калибровки не хватает колонок: " +
               ", ".join(fd.FIELD_LABELS[m] for m in missing_tuning))
else:
    tcol1, tcol2, tcol3 = st.columns(3)
    with tcol1:
        tuning_rho = st.number_input("Плотность нефти (для модели), кг/м³",
                                      min_value=600.0, max_value=1050.0, value=850.0,
                                      step=1.0, key="tuning_rho")
        tuning_api = u.rho_kgm3_to_api(tuning_rho)
        st.caption(f"≈ {tuning_api:.1f} °API — используется как вход в Standing/"
                   f"Beggs-Robinson (в данных месторождения плотность товарной "
                   f"нефти отдельно не измерена, задайте своё лучшее приближение)")
    with tcol2:
        tuning_gamma_g = st.number_input("Отн. плотность газа γg (резервное значение)",
                                          min_value=0.55, max_value=1.70, value=0.80, step=0.01,
                                          key="tuning_gamma_g")
    with tcol3:
        scope_options = ["По всему отфильтрованному массиву"]
        if group_key:
            scope_options.append(f"По каждой группе ({fd.FIELD_LABELS[group_key]})")
        tuning_scope = st.selectbox("Считать коэффициенты", options=scope_options)

    gas_density_available = "gas_rel_density" in mapping and df_filtered["gas_rel_density"].notna().any()
    use_auto_gamma_g = st.checkbox(
        "Брать γg из фактических замеров пробы (среднее по группе), а не из "
        "поля выше — используется колонка «Плотность газа после однократного "
        "разгазирования (по воздуху)», если она сопоставлена",
        value=gas_density_available, disabled=not gas_density_available,
    )
    if not gas_density_available:
        st.caption("В данных нет колонки относительной плотности газа — "
                   "используется резервное значение γg из поля выше для всех групп.")

    tuning_gamma_o = u.rho_kgm3_to_sg(tuning_rho)

    df_for_tuning = df_filtered.copy()
    for pcol in ("pb", "p_sample"):
        if pcol in df_for_tuning.columns:
            df_for_tuning[pcol] = df_for_tuning[pcol].apply(
                lambda v: _to_mpa(v, display_p_unit) if pd.notna(v) else v
            )

    if tuning_scope.startswith("По всему"):
        groups_for_tuning = [("весь массив", df_for_tuning)]
    else:
        groups_for_tuning = [(str(g), gdf) for g, gdf in df_for_tuning.groupby(group_key)]

    tuning_results = {}
    rows_summary = []
    for gname, gdf in groups_for_tuning:
        gamma_g_used = tuning_gamma_g
        gamma_g_source = "вручную"
        if use_auto_gamma_g and "gas_rel_density" in gdf.columns:
            gamma_g_mean = gdf["gas_rel_density"].dropna().mean()
            if pd.notna(gamma_g_mean):
                gamma_g_used = float(gamma_g_mean)
                gamma_g_source = f"факт, n={gdf['gas_rel_density'].notna().sum()}"

        res = pt.tune_group(gdf, tuning_api, gamma_g_used, tuning_gamma_o)
        if res is None:
            continue
        tuning_results[gname] = res
        for prop_key, prop_label in [("pb", "Pb"), ("bo", "Bo"), ("mu", "μo")]:
            f = res[prop_key]
            if f:
                rows_summary.append({
                    "Группа": gname, "Свойство": prop_label, "n": f["n"],
                    "γg использован": f"{gamma_g_used:.3f} ({gamma_g_source})",
                    "Коэффициент": round(f["factor"], 4),
                    "MAPE до, %": round(f["mape_before"], 1),
                    "MAPE после, %": round(f["mape_after"], 1),
                    "R² до": round(f["r2_before"], 3),
                    "R² после": round(f["r2_after"], 3),
                })

    if not rows_summary:
        st.info("Недостаточно данных для расчёта калибровки по выбранным условиям "
                 "(нужно минимум 3 пробы с заполненными Pb/Rs/Bo/μo/T).")
    else:
        st.dataframe(pd.DataFrame(rows_summary), use_container_width=True)

        plot_group_name = st.selectbox("Показать график факт/модель для группы",
                                        options=list(tuning_results.keys()))
        res_plot = tuning_results[plot_group_name]
        df_pred = res_plot["df_pred"]

        plot_specs = [
            ("pb", "pb", "pb_model_mpa", "Pb, МПа"),
            ("bo", "bo", "bo_model", "Bo, д.ед"),
            ("mu", "mu_oil", "mu_model_cp", "μo, мПа·с"),
        ]
        pcols = st.columns(3)
        for col, (prop_key, lab_col, model_col, label) in zip(pcols, plot_specs):
            f = res_plot[prop_key]
            with col:
                if f is None:
                    st.caption(f"{label}: недостаточно данных")
                    continue
                lab_vals = df_pred[lab_col].to_numpy(dtype=float)
                model_vals = df_pred[model_col].to_numpy(dtype=float)
                valid = np.isfinite(lab_vals) & np.isfinite(model_vals)
                lims = [float(np.nanmin(model_vals[valid].tolist() + lab_vals[valid].tolist())),
                        float(np.nanmax(model_vals[valid].tolist() + lab_vals[valid].tolist()))]
                fig_t = go.Figure()
                fig_t.add_trace(go.Scatter(x=model_vals, y=lab_vals, mode="markers",
                                            name="пробы"))
                fig_t.add_trace(go.Scatter(x=lims, y=lims, mode="lines", name="y=x (идеал)",
                                            line=dict(dash="dot", color="gray")))
                fig_t.add_trace(go.Scatter(x=lims, y=[f["factor"] * v for v in lims],
                                            mode="lines", name=f"×{f['factor']:.3f}",
                                            line=dict(color="red")))
                fig_t.update_layout(title=f"{label}: факт vs модель Standing",
                                     xaxis_title="Модель (Standing)", yaxis_title="Факт (лаб.)",
                                     height=350, margin=dict(t=40, b=20))
                st.plotly_chart(fig_t, use_container_width=True)

        st.markdown("**Применить калибровку на главной странице**")
        apply_group = st.selectbox("Коэффициенты какой группы использовать",
                                    options=list(tuning_results.keys()), key="apply_group_select")
        if st.button("Сохранить калибровку и использовать в калькуляторе"):
            res_apply = tuning_results[apply_group]
            st.session_state["pvt_tuning"] = dict(
                source=f"{apply_group} (n={res_apply['n_total']})",
                factor_pb=res_apply["pb"]["factor"] if res_apply["pb"] else 1.0,
                factor_bo=res_apply["bo"]["factor"] if res_apply["bo"] else 1.0,
                factor_mu=res_apply["mu"]["factor"] if res_apply["mu"] else 1.0,
                api=tuning_api, gamma_g=tuning_gamma_g,
            )
            st.success(f"Калибровка сохранена ({apply_group}). Откройте главную "
                       f"страницу — там появится откалиброванная кривая поверх стандартной.")

# ---------------------------------------------------------------------------
# 7. Экспорт очищенной/отфильтрованной таблицы
# ---------------------------------------------------------------------------
st.markdown("---")
st.subheader("7. Экспорт")
export_df = df_filtered.rename(columns={k: _axis_label(k) if k in ("pb", "p_sample")
                                         else f"{fd.FIELD_LABELS.get(k, k)}"
                                         + (f", {fd.FIELD_UNITS[k]}" if fd.FIELD_UNITS.get(k) else "")
                                         for k in df_filtered.columns
                                         if k in fd.FIELD_LABELS})
csv_bytes = export_df.to_csv(index=False).encode("utf-8-sig")
st.download_button("Скачать отфильтрованную таблицу (CSV)", data=csv_bytes,
                    file_name="pvt_field_data_filtered.csv", mime="text/csv")
